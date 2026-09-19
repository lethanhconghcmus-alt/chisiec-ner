"""
audit_date_formula_gr11.py — focused audit cho GR-11 (whole-span DTM cho
date formula hoàn chỉnh). REVIEW ONLY, không tự relabel gì. Bắt buộc quét
các target đã biết (明成化, 莫明德, 莫大正) + quét tổng quát theo pattern
[polity][era][numeral-hoặc-can-chi][marker 年/月/日/科].

Usage:
  python scripts/audit_date_formula_gr11.py \
      --train ... --dev ... --test ... --out-dir artifacts/audit_v1
"""

import argparse
import os
import re
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from src.data_utils import read_conll
from src.audit_utils import extract_all_spans, load_source_map

POLITY_LEXICON = [
    "大越", "大明", "大清", "大元", "大宋", "大唐",
    "明", "莫", "陳", "黎", "李", "晋", "吳", "清", "宋", "元", "唐",
]
NUMERAL_CHARS = "0123456789〇一二三四五六七八九十百千"
CANCHI_CHARS = "甲乙丙丁戊己庚辛壬癸子丑寅卯辰巳午未申酉戌亥"
TIME_MARKER_CHARS = "年月日科"

_polity_re = "|".join(sorted(POLITY_LEXICON, key=len, reverse=True))
CANDIDATE_RE = re.compile(
    rf"(?P<polity>{_polity_re})"
    rf"(?P<era>[一-鿿]{{1,4}}?)"
    rf"(?P<yearpart>[{NUMERAL_CHARS}]+|[{CANCHI_CHARS}]{{2}})"
    rf"(?P<marker>[{TIME_MARKER_CHARS}])"
)

FORCED_TARGETS = ["明成化", "莫明德", "莫大正"]


def find_candidates_in_text(text: str) -> list:
    out = []
    seen_spans = set()
    for m in CANDIDATE_RE.finditer(text):
        span = (m.start(), m.end() - 1)
        if span in seen_spans:
            continue
        seen_spans.add(span)
        out.append({
            "polity": m.group("polity"), "era": m.group("era"),
            "yearpart": m.group("yearpart"), "marker": m.group("marker"),
            "start": m.start(), "end": m.end() - 1, "full_match": m.group(0),
            "time_marker_evidence": m.group("yearpart") + m.group("marker"),
        })
    # Bo sung cac target bat buoc neu regex chinh chua bat duoc het bien the
    for target in FORCED_TARGETS:
        idx = 0
        while True:
            pos = text.find(target, idx)
            if pos == -1:
                break
            span = (pos, pos + len(target) - 1)
            if span not in seen_spans:
                # tim marker/numeral/canchi ngay sau target neu co
                trailing = text[pos + len(target): pos + len(target) + 6]
                m2 = re.match(rf"([{NUMERAL_CHARS}]+|[{CANCHI_CHARS}]{{2}})([{TIME_MARKER_CHARS}])", trailing)
                if m2:
                    full_end = pos + len(target) + m2.end() - 1
                    seen_spans.add((pos, full_end))
                    out.append({
                        "polity": target[0], "era": target[1:], "yearpart": m2.group(1),
                        "marker": m2.group(2), "start": pos, "end": full_end,
                        "full_match": text[pos:full_end + 1],
                        "time_marker_evidence": m2.group(0),
                        "forced_target": target,
                    })
                else:
                    seen_spans.add(span)
                    out.append({
                        "polity": target[0], "era": target[1:], "yearpart": "", "marker": "",
                        "start": pos, "end": span[1], "full_match": target,
                        "time_marker_evidence": "(khong tim thay numeral/can-chi+marker ngay sau)",
                        "forced_target": target,
                    })
            idx = pos + 1
    return out


def classify_gold_configuration(cand: dict, entities_here: list) -> tuple:
    """Trả (config_label, overlapping_entities_desc)."""
    cs, ce = cand["start"], cand["end"]
    overlapping = [e for e in entities_here if not (e["end"] < cs or e["start"] > ce)]
    desc = "; ".join(f"{e['surface']}/{e['label']}[{e['start']}-{e['end']}]" for e in overlapping)

    if not overlapping:
        return "unannotated", desc

    exact = [e for e in overlapping if e["start"] == cs and e["end"] == ce]
    if len(exact) == 1:
        label = exact[0]["label"]
        return ("whole_DTM" if label == "DTM" else label), desc

    covers_fully = (min(e["start"] for e in overlapping) == cs and max(e["end"] for e in overlapping) == ce)
    sorted_ov = sorted(overlapping, key=lambda e: e["start"])
    contiguous = all(
        b["start"] == a["end"] + 1 for a, b in zip(sorted_ov, sorted_ov[1:])
    ) if len(sorted_ov) > 1 else True

    if covers_fully and contiguous and len(sorted_ov) >= 2:
        if all(e["label"] == "DTM" for e in sorted_ov):
            return "split_DTM", desc
        return "mixed/overlap", desc

    return "mixed/overlap", desc


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--train", required=True)
    ap.add_argument("--dev", required=True)
    ap.add_argument("--test", required=True)
    ap.add_argument("--source-map", default=None)
    ap.add_argument("--out-dir", default="artifacts/audit_v1")
    args = ap.parse_args()

    splits_data = {
        "train": read_conll(args.train),
        "dev": read_conll(args.dev),
        "test": read_conll(args.test),
    }
    source_map = load_source_map(args.source_map)
    entities = extract_all_spans(splits_data, source_map, context_window=30)

    entities_by_sentence = defaultdict(list)
    for e in entities:
        entities_by_sentence[(e["split"], e["sample_id"])].append(e)

    rows = []
    candidate_id = 0
    for split, sentences in splits_data.items():
        for idx, (tokens, labels) in enumerate(sentences):
            text = "".join(tokens)
            candidates = find_candidates_in_text(text)
            if not candidates:
                continue
            ents_here = entities_by_sentence.get((split, idx), [])
            doc_id = ents_here[0].get("document_id") if ents_here else None
            for cand in candidates:
                config, overlap_desc = classify_gold_configuration(cand, ents_here)
                ctx_s, ctx_e = max(0, cand["start"] - 30), min(len(text), cand["end"] + 1 + 30)
                marked = f"{text[ctx_s:cand['start']]}【{cand['full_match']}】{text[cand['end']+1:ctx_e]}"

                if config == "unannotated":
                    confidence = "low"
                elif config in ("whole_DTM", "split_DTM"):
                    confidence = "high"
                elif config == "mixed/overlap":
                    confidence = "medium"
                else:
                    confidence = "medium"

                rows.append({
                    "candidate_id": candidate_id,
                    "split": split, "sample_id": idx, "document_id": doc_id,
                    "full_text": text, "marked_context": marked,
                    "original_entities": overlap_desc,
                    "candidate_span": f"{cand['start']}-{cand['end']}",
                    "candidate_surface": cand["full_match"],
                    "current_gold_configuration": config,
                    "time_marker_evidence": cand["time_marker_evidence"],
                    "forced_target": cand.get("forced_target", ""),
                    "candidate_confidence": confidence,
                    "reviewer_decision": "",
                    "final_span": "",
                    "final_label": "",
                    "suggested_rule_id": "GR-11",
                    "reviewer_notes": "",
                })
                candidate_id += 1

    review_dir = os.path.join(args.out_dir, "review")
    reports_dir = os.path.join(args.out_dir, "reports")
    os.makedirs(review_dir, exist_ok=True)
    os.makedirs(reports_dir, exist_ok=True)

    df = pd.DataFrame(rows)
    out_xlsx = os.path.join(review_dir, "gr11_date_formula_candidates.xlsx")
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="gr11_candidates", index=False)
        ws = writer.sheets["gr11_candidates"]
        ws.freeze_panes = "A2"
        wrap_cols = {"full_text", "marked_context", "original_entities"}
        for col_idx, col_name in enumerate(df.columns, start=1):
            letter = get_column_letter(col_idx)
            if col_name in wrap_cols:
                ws.column_dimensions[letter].width = 60
                for row in range(2, len(df) + 2):
                    ws.cell(row=row, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
            else:
                ws.column_dimensions[letter].width = max(len(col_name) + 2, 12)

    # ── Summary report ───────────────────────────────────────────────────
    lines = ["# GR-11 focused audit — date formula candidates", ""]
    lines.append(f"Tổng candidate: **{len(rows)}**")
    lines.append("")
    lines.append("**Không khẳng định mọi candidate là lỗi** — bảng dưới chỉ mô tả gold hiện tại.")
    lines.append("")
    lines.append("## Count theo pattern (polity prefix)")
    by_polity = defaultdict(int)
    for r in rows:
        by_polity[r["candidate_surface"][:1] if not r["forced_target"] else r["forced_target"][0]] += 1
    for polity, c in sorted(by_polity.items(), key=lambda kv: -kv[1]):
        lines.append(f"- {polity}: {c}")
    lines.append("")
    lines.append("## Count theo current_gold_configuration")
    by_config = defaultdict(int)
    for r in rows:
        by_config[r["current_gold_configuration"]] += 1
    for config, c in sorted(by_config.items(), key=lambda kv: -kv[1]):
        lines.append(f"- {config}: {c}")
    lines.append("")
    lines.append("## Toàn bộ occurrence 莫明德 / 莫大正")
    for r in rows:
        if r["forced_target"] in ("莫明德", "莫大正"):
            lines.append(f"- [{r['split']}:{r['sample_id']}] config={r['current_gold_configuration']} | {r['marked_context']}")
    lines.append("")
    n_support_whole = by_config.get("whole_DTM", 0)
    n_needs_exception = sum(v for k, v in by_config.items() if k not in ("whole_DTM", "unannotated"))
    lines.append(f"## Support whole-span DTM vs cần exception")
    lines.append(f"- Đã là whole_DTM (khớp policy mới, không cần sửa): {n_support_whole}")
    lines.append(f"- Cấu hình khác (split_DTM/PER/ORG/LOC/mixed — cần xem từng case, "
                 f"KHÔNG mặc định coi là lỗi): {n_needs_exception}")
    lines.append(f"- Chưa có nhãn (unannotated): {by_config.get('unannotated', 0)}")

    with open(os.path.join(reports_dir, "gr11_date_formula_summary.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print(f"GR-11 candidates: {len(rows)}")
    print(f"By config: {dict(by_config)}")
    print(f"Output -> {out_xlsx}, {reports_dir}/gr11_date_formula_summary.md")
    print("\n=== 莫明德 / 莫大正 occurrences ===")
    for r in rows:
        if r["forced_target"] in ("莫明德", "莫大正"):
            print(f"  [{r['split']}:{r['sample_id']}] config={r['current_gold_configuration']} {r['marked_context']}")


if __name__ == "__main__":
    main()
