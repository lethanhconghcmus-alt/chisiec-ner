"""
export_yushi_gudu_transaction.py — tạo transaction review riêng cho
train:437 (mục B): REMOVE_ENTITY(故都/LOC) + CORRECT_BOUNDARY(御史->都御史).
KHÔNG prefill quyết định cuối, KHÔNG tự chạy transaction này cho tới khi
reviewer điền đủ field bắt buộc + 3 checkbox xác nhận.

Usage:
  python scripts/export_yushi_gudu_transaction.py \
      --train ... --dev ... --test ... \
      --out artifacts/audit_v1/review/yushi_gudu_transaction.xlsx
"""

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.data_utils import read_conll
from src.audit_utils import extract_all_spans, load_source_map
from src.adjudication import TRANSACTION_ACTION_OPTIONS
from src.normalization import codepoint_category


SPLIT, SAMPLE_ID = "train", 437
REMOVE_ENTITY_ID = 6272   # 故都/LOC
BOUNDARY_ENTITY_ID = 6273  # 御史/TITLE


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--train", required=True)
    ap.add_argument("--dev", required=True)
    ap.add_argument("--test", required=True)
    ap.add_argument("--source-map", default=None)
    ap.add_argument("--out", default="artifacts/audit_v1/review/yushi_gudu_transaction.xlsx")
    args = ap.parse_args()

    splits_data = {
        "train": read_conll(args.train),
        "dev": read_conll(args.dev),
        "test": read_conll(args.test),
    }
    source_map = load_source_map(args.source_map)
    entities = extract_all_spans(splits_data, source_map, context_window=80)
    entity_by_id = {e["entity_id"]: e for e in entities}

    e_remove = entity_by_id.get(REMOVE_ENTITY_ID)
    e_boundary = entity_by_id.get(BOUNDARY_ENTITY_ID)
    if e_remove is None or e_boundary is None:
        print(f"WARNING: khong tim thay entity_id={REMOVE_ENTITY_ID} hoac {BOUNDARY_ENTITY_ID} "
              f"trong dataset hien tai -- co the da doi so voi luc audit truoc. Kiem tra lai thu cong.")
        return

    text = e_remove["text"]
    ents_here = sorted([e for e in entities if e["split"] == SPLIT and e["sample_id"] == SAMPLE_ID],
                        key=lambda e: e["start"])
    original_entities_str = "; ".join(
        f"{e['surface']}/{e['label']}[{e['start']}-{e['end']}](id={e['entity_id']})" for e in ents_here
    )

    ctx_s, ctx_e = max(0, e_remove["start"] - 80), min(len(text), e_boundary["end"] + 1 + 80)
    marked = (
        f"{text[ctx_s:e_remove['start']]}"
        f"《{text[e_remove['start']:e_remove['end']+1]}》"
        f"{text[e_remove['end']+1:e_boundary['start']]}"
        f"〖{text[e_boundary['start']:e_boundary['end']+1]}〗"
        f"{text[e_boundary['end']+1:ctx_e]}"
    )

    pua_chars_in_range = [
        (i, ch, codepoint_category(ord(ch)))
        for i, ch in enumerate(text[e_remove["start"]:e_boundary["end"] + 1], start=e_remove["start"])
        if codepoint_category(ord(ch)) is not None
    ]
    pua_status = (
        "KHONG co ky tu PUA trong pham vi 2 entity nay"
        if not pua_chars_in_range else
        f"CO {len(pua_chars_in_range)} ky tu nghi van: {pua_chars_in_range}"
    )

    rows = [
        {
            "transaction_id": "T-yushi-gudu-437",
            "action": "REMOVE_ENTITY",
            "split": SPLIT, "sample_id": SAMPLE_ID, "document_id": e_remove.get("document_id"),
            "full_text": text,
            "marked_context_pm80": marked,
            "raw_text_segment": text[e_remove["start"]:e_remove["end"] + 1],
            "normalized_text_segment": text[e_remove["start"]:e_remove["end"] + 1],  # khong co PUA trong doan nay theo scan
            "pua_normalization_status": pua_status,
            "original_entities_in_sentence": original_entities_str,
            "source_entity_id": REMOVE_ENTITY_ID,
            "old_label": e_remove["label"], "old_start": e_remove["start"], "old_end": e_remove["end"],
            "old_surface": e_remove["surface"],
            "final_label": "", "final_start": "", "final_end": "",
            "remove_reason": "",
            "guideline_rule_id": "",
            "reviewer_notes": "",
            "confirm_a_gudu_not_valid_LOC": "",
            "confirm_b_duyushi_is_valid_full_TITLE": "",
            "confirm_c_flat_ner_priority_approved": "",
        },
        {
            "transaction_id": "T-yushi-gudu-437",
            "action": "CORRECT_BOUNDARY",
            "split": SPLIT, "sample_id": SAMPLE_ID, "document_id": e_boundary.get("document_id"),
            "full_text": text,
            "marked_context_pm80": marked,
            "raw_text_segment": text[e_boundary["start"]:e_boundary["end"] + 1],
            "normalized_text_segment": text[e_boundary["start"]:e_boundary["end"] + 1],
            "pua_normalization_status": pua_status,
            "original_entities_in_sentence": original_entities_str,
            "source_entity_id": BOUNDARY_ENTITY_ID,
            "old_label": e_boundary["label"], "old_start": e_boundary["start"], "old_end": e_boundary["end"],
            "old_surface": e_boundary["surface"],
            "final_label": "", "final_start": "", "final_end": "",
            "remove_reason": "",
            "guideline_rule_id": "",
            "reviewer_notes": "",
            "confirm_a_gudu_not_valid_LOC": "",
            "confirm_b_duyushi_is_valid_full_TITLE": "",
            "confirm_c_flat_ner_priority_approved": "",
        },
    ]

    df = pd.DataFrame(rows)
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with pd.ExcelWriter(args.out, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="yushi_gudu_transaction", index=False)
        ws = writer.sheets["yushi_gudu_transaction"]
        ws.freeze_panes = "A2"
        wrap_cols = {"full_text", "marked_context_pm80", "original_entities_in_sentence", "reviewer_notes"}
        for col_idx, col_name in enumerate(df.columns, start=1):
            letter = get_column_letter(col_idx)
            if col_name in wrap_cols:
                ws.column_dimensions[letter].width = 60
                for row in range(2, len(df) + 2):
                    ws.cell(row=row, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
            else:
                ws.column_dimensions[letter].width = max(len(col_name) + 2, 14)

        for col_name, options in {
            "action": sorted(TRANSACTION_ACTION_OPTIONS),
            "confirm_a_gudu_not_valid_LOC": ["YES", "NO"],
            "confirm_b_duyushi_is_valid_full_TITLE": ["YES", "NO"],
            "confirm_c_flat_ner_priority_approved": ["YES", "NO"],
        }.items():
            col_idx = list(df.columns).index(col_name) + 1
            letter = get_column_letter(col_idx)
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{len(df) + 1}")

    print(f"Exported transaction T-yushi-gudu-437 (2 action) -> {args.out}")
    print(f"Context: {marked}")
    print(f"PUA status: {pua_status}")
    print("KHONG tu chay transaction nay -- cho reviewer dien du field bat buoc + 3 checkbox xac nhan.")


if __name__ == "__main__":
    main()
