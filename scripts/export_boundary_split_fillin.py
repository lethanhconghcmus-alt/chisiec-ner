"""
export_boundary_split_fillin.py — xuất đúng các case CORRECT_BOUNDARY/
SPLIT_ENTITY đang thiếu final_start/final_end cấu trúc (blocker phát hiện
bởi apply_adjudication.py --dry-run), để reviewer điền offset mới tường
minh. Đọc-only, không sửa gì.

Usage:
  python scripts/export_boundary_split_fillin.py \
      --train ... --dev ... --test ... \
      --review-dir artifacts/audit_v1/review \
      --out artifacts/audit_v1/review/boundary_split_fill_in.xlsx
"""

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from src.data_utils import read_conll
from src.audit_utils import extract_all_spans, load_source_map


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--train", required=True)
    ap.add_argument("--dev", required=True)
    ap.add_argument("--test", required=True)
    ap.add_argument("--source-map", default=None)
    ap.add_argument("--review-dir", required=True)
    ap.add_argument("--out", default="artifacts/audit_v1/review/boundary_split_fill_in.xlsx")
    args = ap.parse_args()

    splits_data = {
        "train": read_conll(args.train),
        "dev": read_conll(args.dev),
        "test": read_conll(args.test),
    }
    source_map = load_source_map(args.source_map)
    entities = extract_all_spans(splits_data, source_map, context_window=30)
    entity_by_id = {e["entity_id"]: e for e in entities}

    rows = []

    # ── 1. deepdive_occurrences: CORRECT_BOUNDARY / SPLIT_ENTITY ────────
    guideline_path = os.path.join(args.review_dir, "review_guideline_ambiguities.xlsx")
    df_deep = pd.read_excel(guideline_path, sheet_name="deepdive_occurrences")
    for _, r in df_deep[df_deep["reviewer_decision"].isin(["CORRECT_BOUNDARY", "SPLIT_ENTITY"])].iterrows():
        e = entity_by_id.get(r["entity_id"])
        rows.append({
            "row_type": r["reviewer_decision"],
            "source": "deepdive_occurrences",
            "entity_id": r["entity_id"],
            "surface_hint": r["surface"],
            "split": r["split"],
            "sample_id": r["sample_id"],
            "document_id": r["document_id"],
            "gold_surface": r["gold_surface"],
            "gold_start": r["gold_start"],
            "gold_end": r["gold_end"],
            "gold_label": r["gold_label"],
            "full_sentence": e["text"] if e else r["full_sentence"],
            "context_pm30": e["marked_context"] if e else r["context_pm30"],
            "existing_reviewer_notes": "",  # deepdive rows khong co reviewer_notes rieng, xem guideline_ambiguities note cho ca surface
            "final_label": "",
            "final_start": "",
            "final_end": "",
            "split2_start": "",
            "split2_end": "",
            "split2_label": "",
            "split3_start_optional": "",
            "split3_end_optional": "",
            "split3_label_optional": "",
            "fill_notes": "",
        })

    # ── 2. singleton_anomalies: 大行 (CORRECT_BOUNDARY, thiếu final_label lẫn offset) ──
    singleton_path = os.path.join(args.review_dir, "review_priority_singleton_anomalies.xlsx")
    df_single = pd.read_excel(singleton_path)
    for _, r in df_single[df_single["reviewer_decision"] == "CORRECT_BOUNDARY"].iterrows():
        ref = str(r["record_ref"])
        split, sample_id = ref.split(":", 1)
        sample_id = int(sample_id)
        # tim entity tuong ung qua CA surface LAN minority_label trong cau do
        # (BUG da fix: chi loc theo label la khong du khi 1 cau co nhieu
        # entity cung nhan, vd 3 TITLE trong 1 cau -- se lay nham entity
        # dau tien thay vi dung surface, xem case 大行/經畧使 test:139).
        cand = [e for e in entities if e["split"] == split and e["sample_id"] == sample_id
                and e["label"] == r["minority_label"] and e["surface"] == r["surface"]]
        if len(cand) > 1:
            print(f"WARNING: {len(cand)} entity trung surface={r['surface']!r} "
                  f"label={r['minority_label']!r} tai {split}:{sample_id} -- lay entity dau tien, "
                  f"KIEM TRA LAI thu cong.")
        e = cand[0] if cand else None
        if e is None:
            print(f"WARNING: KHONG tim thay entity khop surface={r['surface']!r} "
                  f"label={r['minority_label']!r} tai {split}:{sample_id} -- dong nay se thieu du lieu song.")
        rows.append({
            "row_type": "CORRECT_BOUNDARY",
            "source": "singleton_anomalies",
            "entity_id": e["entity_id"] if e else "",
            "surface_hint": r["surface"],
            "split": split,
            "sample_id": sample_id,
            "document_id": e.get("document_id") if e else r.get("minority_document_id"),
            "gold_surface": e["surface"] if e else r["surface"],
            "gold_start": e["start"] if e else "",
            "gold_end": e["end"] if e else "",
            "gold_label": e["label"] if e else r["minority_label"],
            "full_sentence": e["text"] if e else "",
            "context_pm30": e["marked_context"] if e else r["minority_context"],
            "existing_reviewer_notes": r["reviewer_notes"],
            "final_label": "",
            "final_start": "",
            "final_end": "",
            "split2_start": "",
            "split2_end": "",
            "split2_label": "",
            "split3_start_optional": "",
            "split3_end_optional": "",
            "split3_label_optional": "",
            "fill_notes": "",
        })

    df_out = pd.DataFrame(rows)
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with pd.ExcelWriter(args.out, engine="openpyxl") as writer:
        df_out.to_excel(writer, sheet_name="boundary_split_fillin", index=False)
        ws = writer.sheets["boundary_split_fillin"]
        ws.freeze_panes = "A2"
        wrap_cols = {"full_sentence", "context_pm30", "existing_reviewer_notes", "fill_notes"}
        for col_idx, col_name in enumerate(df_out.columns, start=1):
            letter = get_column_letter(col_idx)
            if col_name in wrap_cols:
                ws.column_dimensions[letter].width = 55
                for row in range(2, len(df_out) + 2):
                    ws.cell(row=row, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
            else:
                ws.column_dimensions[letter].width = max(len(col_name) + 2, 12)

    n_boundary = sum(1 for r in rows if r["row_type"] == "CORRECT_BOUNDARY")
    n_split = sum(1 for r in rows if r["row_type"] == "SPLIT_ENTITY")
    print(f"Exported {len(rows)} row(s) -> {args.out}")
    print(f"  CORRECT_BOUNDARY: {n_boundary}")
    print(f"  SPLIT_ENTITY: {n_split}")


if __name__ == "__main__":
    main()
